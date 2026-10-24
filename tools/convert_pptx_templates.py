"""Converte os PPTX de antes e depois (exportados do Canva) em modelos do bot.

Para cada slide: separa o que é FIXO (background.png: logo, círculos, títulos, rodapé)
do que é VARIÁVEL (fotos e textos, configurados no template.json) e do que fica NA FRENTE
da foto (overlay.png: moldura branca e selo ANTES/DEPOIS).

Requer LibreOffice (soffice) e pdftoppm. Uso:
    python tools/convert_pptx_templates.py <pasta_com_os_pptx>
"""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]
EMU_PX = 6350  # 6.858.000 EMU = 1080 px

# o que é variável / vai por cima da foto, por slide (nomes dos shapes no PPTX)
SPECS = {
    "story-antes-e-depois": [{
        "id": "story_antes_depois", "name": "Story – Antes e depois", "desc": "Story com duas fotos lado a lado (antes e depois) e o serviço realizado.", "size": (1080, 1920),
        "photos": [("antes", "Foto ANTES", "Image 1"), ("depois", "Foto DEPOIS", "Image 2")],
        "texts": [("servico", "Qual serviço foi feito? (ex.: Corte)", "Text 3", 20, "Poppins-Regular.ttf", None, "#9E2160")],
        "overlay": ["Shape 4", "Text 5", "Shape 6", "Text 7"],
    }],
    "feed-carrossel-antes-e-depois": [
        {"id": "feed_antes", "hidden": True, "desc": "Slide do antes (uso interno do carrossel).", "name": "Feed – Antes", "size": (1080, 1350),
         "photos": [("foto", "Foto ANTES", "Image 1")],
         "texts": [("servico", "Qual serviço foi feito? (ex.: Corte)", "Text 4", 25, "Lora-Variable.ttf", "Bold", "#4D1439")],
         "overlay": ["Shape 2", "Text 3"]},
        {"id": "feed_depois", "hidden": True, "desc": "Slide do depois (uso interno do carrossel).", "name": "Feed – Depois", "size": (1080, 1350),
         "photos": [("foto", "Foto DEPOIS", "Image 1")],
         "texts": [("servico", "Qual serviço foi feito? (ex.: Corte)", "Text 4", 25, "Lora-Variable.ttf", "Bold", "#4D1439")],
         "overlay": ["Shape 2", "Text 3"]},
        None,  # slide 3 é 100% fixo (CTA): exportado como arte estática
    ],
}


def render(prs_path: Path, out_prefix: Path, size: tuple[int, int]) -> list[Path]:
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out_prefix.parent), str(prs_path)],
                   check=True, capture_output=True)
    pdf = out_prefix.parent / (prs_path.stem + ".pdf")
    subprocess.run(["pdftoppm", "-png", "-scale-to-x", str(size[0]), "-scale-to-y", str(size[1]), str(pdf), str(out_prefix)],
                   check=True)
    return sorted(out_prefix.parent.glob(out_prefix.name + "-*.png"))


def drop(slide, names: set[str], keep: bool = False) -> None:
    for sh in list(slide.shapes):
        if (sh.name in names) != keep:
            sh._element.getparent().remove(sh._element)


def rect(sh) -> dict:
    return {"x": round(sh.left / EMU_PX), "y": round(sh.top / EMU_PX),
            "width": round(sh.width / EMU_PX), "height": round(sh.height / EMU_PX)}


def alpha_from_two_renders(white: Path, black: Path) -> Image.Image:
    w = np.asarray(Image.open(white).convert("RGB"), dtype=np.float32)
    k = np.asarray(Image.open(black).convert("RGB"), dtype=np.float32)
    a = np.clip(1 - (w - k).mean(axis=2) / 255.0, 0, 1)
    rgb = np.where(a[..., None] > 0.003, k / np.maximum(a[..., None], 0.003), 0)
    out = np.dstack([np.clip(rgb, 0, 255), a * 255]).astype(np.uint8)
    return Image.fromarray(out, "RGBA")


def measure_frame(media: Path, slot_w: int) -> tuple[int, int]:
    """Largura da borda branca e raio dos cantos, medidos na foto de exemplo do PPTX e levados ao canvas."""
    im = Image.open(media).convert("RGBA")
    scale = slot_w / im.width
    mid = im.height // 2
    b = 0
    while b < im.width and im.getpixel((b, mid))[:3] == (255, 255, 255):
        b += 1
    alpha = im.getchannel("A")
    best = min(range(5, 200), key=lambda r: np.abs(
        np.asarray(alpha, dtype=np.int16) - np.asarray(_rr((im.width, im.height), r), dtype=np.int16)).sum())
    return max(1, round(b * scale)), round(best * scale)


def _rr(size: tuple[int, int], r: int) -> Image.Image:
    m = Image.new("L", size, 0)
    ImageDraw.Draw(m).rounded_rectangle((0, 0, size[0] - 1, size[1] - 1), radius=r, fill=255)
    return m


def ring(size: tuple[int, int], box: dict, b: int, r: int) -> Image.Image:
    s = 4
    big = Image.new("L", (size[0] * s, size[1] * s), 0)
    d = ImageDraw.Draw(big)
    x0, y0, x1, y1 = box["x"] * s, box["y"] * s, (box["x"] + box["width"]) * s - 1, (box["y"] + box["height"]) * s - 1
    d.rounded_rectangle((x0, y0, x1, y1), radius=r * s, fill=255)
    d.rounded_rectangle((x0 + b * s, y0 + b * s, x1 - b * s, y1 - b * s), radius=max(1, (r - b)) * s, fill=0)
    mask = big.resize(size, Image.LANCZOS)
    out = Image.new("RGBA", size, (255, 255, 255, 0))
    out.putalpha(mask)
    return out


def main(src: Path) -> None:
    tmp = Path(tempfile.mkdtemp())
    for stem, slides in SPECS.items():
        pptx = src / f"{stem}.pptx"
        media_dir = tmp / f"{stem}_x"
        subprocess.run(["unzip", "-o", "-q", str(pptx), "-d", str(media_dir)], check=True)
        size = next(s["size"] for s in slides if s)
        ref = Presentation(str(pptx))

        # 1) fundo: remove variáveis e overlay
        bgp = Presentation(str(pptx))
        for sl, spec in zip(bgp.slides, slides):
            if spec:
                drop(sl, {p[2] for p in spec["photos"]} | {t[2] for t in spec["texts"]} | set(spec["overlay"]))
        bg_in = tmp / f"{stem}_bg.pptx"; bgp.save(bg_in)
        bgs = render(bg_in, tmp / f"{stem}_bg", size)

        # 2) overlay: só os elementos de frente, sobre branco e sobre preto
        overlays = {}
        for tone, rgb in (("w", RGBColor(255, 255, 255)), ("k", RGBColor(0, 0, 0))):
            p = Presentation(str(pptx))
            for sl, spec in zip(p.slides, slides):
                if spec:
                    drop(sl, set(spec["overlay"]), keep=True)
                sl.background.fill.solid(); sl.background.fill.fore_color.rgb = rgb
            f = tmp / f"{stem}_ov{tone}.pptx"; p.save(f)
            overlays[tone] = render(f, tmp / f"{stem}_ov{tone}", size)

        for i, spec in enumerate(slides):
            if spec is None:
                out = ROOT / "templates" / "feed_antes_depois"; out.mkdir(parents=True, exist_ok=True)
                # slide 3 original, completo (todos os textos são fixos)
                full = tmp / f"{stem}_full"
                full_png = render(pptx, full, size)[i]
                Image.open(full_png).convert("RGB").save(out / "cta.png")
                continue
            d = ROOT / "templates" / spec["id"]; d.mkdir(parents=True, exist_ok=True)
            Image.open(bgs[i]).convert("RGB").save(d / "background.png")
            ov = alpha_from_two_renders(overlays["w"][i], overlays["k"][i])
            shapes = {sh.name: sh for sh in ref.slides[i].shapes}
            photos, texts = [], []
            for pid, label, shape in spec["photos"]:
                box = rect(shapes[shape])
                # foto de exemplo do PPTX (só para medir borda e raio): a de mesma proporção do slot
                cand = [m for m in (media_dir / "ppt" / "media").glob(f"image-{i+1}-*.png")
                        if abs(Image.open(m).width / Image.open(m).height - box["width"] / box["height"]) < 0.02
                        and Image.open(m).width > 600]
                b, r = measure_frame(cand[0], box["width"])
                ov = Image.alpha_composite(ring(size, box, b, r), ov)
                photos.append({"id": pid, "label": label, **box, "fit": "cover", "mask": "rounded", "radius": r})
            for tid, label, shape, pt, font, var, color in spec["texts"]:
                t = {"id": tid, "label": label, **rect(shapes[shape]), "max_font_size": pt * 2, "min_font_size": round(pt * 2 * 0.6),
                     "color": color, "font": font, "align": "center", "valign": "middle", "max_lines": 1, "max_chars": 40}
                if var:
                    t["font_variation"] = var
                texts.append(t)
            ov.save(d / "overlay.png")
            cfg = {"id": spec["id"], "name": spec["name"],
                   "description": spec.get("desc", ""),
                   "canvas": {"width": size[0], "height": size[1]}, "background": "background.png",
                   "overlay": "overlay.png", "photos": photos, "texts": texts}
            if spec.get("hidden"):
                cfg["hidden"] = True
            (d / "template.json").write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
            print("OK", d.name, [p["radius"] for p in photos])


def write_carousel() -> None:
    """Carrossel do feed: uma única opção no menu (pede ANTES e DEPOIS) que gera 3 slides."""
    d = ROOT / "templates" / "feed_antes_depois"
    cfg = {
        "type": "carousel", "id": "feed_antes_depois", "name": "Feed – Antes e depois (carrossel)",
        "description": "Carrossel de 3 imagens: antes, depois e o convite para agendar.",
        "photos": [{"id": "antes", "label": "Foto ANTES"}, {"id": "depois", "label": "Foto DEPOIS"}],
        "texts": [{"id": "servico", "label": "Qual serviço foi feito? (ex.: Corte)", "max_chars": 40}],
        "slides": [
            {"template": "feed_antes", "photos": {"foto": "antes"}, "texts": {"servico": "servico"}},
            {"template": "feed_depois", "photos": {"foto": "depois"}, "texts": {"servico": "servico"}},
            {"static": "cta.png"},
        ],
    }
    (d / "template.json").write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main(Path(sys.argv[1]))
    write_carousel()
